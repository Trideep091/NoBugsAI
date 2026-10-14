"""Excel Workbook Report Generation Service (Milestone 9)"""
import io
import json
from typing import List, Dict, Any
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from app.db.models import IncidentModel, LogEventModel, LogFileModel
from app.pipeline.redactor import redactor


class ExcelReportService:
    """Generates formatted multi-sheet Excel reports with styling and formula injection defense"""

    def generate_incident_report(
        self,
        incidents: List[IncidentModel],
        events: List[LogEventModel],
        files: List[LogFileModel],
    ) -> io.BytesIO:
        """
        Creates a formatted 3-worksheet .xlsx workbook:
        Sheet 1: Incident Summary
        Sheet 2: Event Timeline
        Sheet 3: Original Log Events
        """
        wb = Workbook()

        # Styles
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        title_font = Font(name="Calibri", size=14, bold=True, color="0F172A")
        sub_font = Font(name="Calibri", size=10, italic=True, color="64748B")
        regular_font = Font(name="Calibri", size=10)
        bold_font = Font(name="Calibri", size=10, bold=True)
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )

        # Severity fills
        sev_fills = {
            "Critical": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
            "High": PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid"),
            "Mild": PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid"),
        }
        sev_fonts = {
            "Critical": Font(name="Calibri", size=10, bold=True, color="991B1B"),
            "High": Font(name="Calibri", size=10, bold=True, color="9A3412"),
            "Mild": Font(name="Calibri", size=10, bold=True, color="854D0E"),
        }

        # Filename lookup
        file_name_map = {f.id: f.filename for f in files}

        # -------------------------------------------------------------
        # Worksheet 1: Incident Summary
        # -------------------------------------------------------------
        ws1 = wb.active
        ws1.title = "Incident Summary"
        ws1.views.sheetView[0].showGridLines = True

        # Sheet Title block
        ws1["A1"] = "NoBugs AI — Incident Investigation Summary Report"
        ws1["A1"].font = title_font
        ws1["A2"] = "Generated from sanitized log analysis and evidence-backed root cause diagnostic pipeline"
        ws1["A2"].font = sub_font

        headers1 = [
            "Incident ID",
            "Title",
            "Severity",
            "Lifecycle Status",
            "First Seen",
            "Last Seen",
            "Occurrences",
            "Root-Cause Hypothesis",
            "AI-Recommended Solution",
            "Confidence Assessment",
            "Verification Steps",
        ]

        row_idx = 4
        for col_idx, h in enumerate(headers1, start=1):
            cell = ws1.cell(row=row_idx, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for inc in incidents:
            row_idx += 1
            # Format recommended fixes
            fixes_text = ""
            if inc.recommended_fixes:
                fixes_text = "; ".join([
                    f"{idx+1}. {rf.get('action', '')} (Reason: {rf.get('reason', '')})"
                    for idx, rf in enumerate(inc.recommended_fixes)
                ])

            verif_text = "; ".join(inc.verification_steps or [])

            vals = [
                redactor.sanitize_for_spreadsheet(inc.id[:8]),
                redactor.sanitize_for_spreadsheet(inc.title),
                inc.severity,
                inc.lifecycle_status,
                inc.first_seen.strftime("%Y-%m-%d %H:%M:%S") if inc.first_seen else "N/A",
                inc.last_seen.strftime("%Y-%m-%d %H:%M:%S") if inc.last_seen else "N/A",
                inc.occurrence_count,
                redactor.sanitize_for_spreadsheet(inc.root_cause_hypothesis or "Pending Analysis"),
                redactor.sanitize_for_spreadsheet(fixes_text or "N/A"),
                redactor.sanitize_for_spreadsheet(f"{inc.ai_confidence or 'Medium'} ({inc.ai_confidence_rationale or 'Telemetry derived'})"),
                redactor.sanitize_for_spreadsheet(verif_text or "N/A"),
            ]

            for col_idx, val in enumerate(vals, start=1):
                cell = ws1.cell(row=row_idx, column=col_idx, value=val)
                cell.font = regular_font
                cell.border = thin_border
                cell.alignment = Alignment(vertical="top", wrap_text=(col_idx in (2, 8, 9, 10, 11)))

                # Severity styling
                if col_idx == 3 and val in sev_fills:
                    cell.fill = sev_fills[val]
                    cell.font = sev_fonts[val]
                    cell.alignment = Alignment(horizontal="center", vertical="top")

        # -------------------------------------------------------------
        # Worksheet 2: Event Timeline
        # -------------------------------------------------------------
        ws2 = wb.create_sheet(title="Event Timeline")
        ws2.views.sheetView[0].showGridLines = True

        ws2["A1"] = "NoBugs AI — Event Timeline"
        ws2["A1"].font = title_font
        ws2["A2"] = "Chronological sequence of correlated error events and downstream symptoms"
        ws2["A2"].font = sub_font

        headers2 = [
            "Event ID",
            "Timestamp",
            "Incident ID",
            "Source Filename",
            "Original Line Number",
            "Component",
            "Sanitized Message",
            "Event Classification",
        ]

        row_idx = 4
        for col_idx, h in enumerate(headers2, start=1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Only error events or correlated events for timeline
        timeline_events = [e for e in events if e.is_error or e.incident_id]
        timeline_events.sort(key=lambda e: e.timestamp if e.timestamp else e.original_line_number)

        for ev in timeline_events:
            row_idx += 1
            vals = [
                redactor.sanitize_for_spreadsheet(ev.id[:8]),
                ev.timestamp.strftime("%Y-%m-%d %H:%M:%S") if ev.timestamp else "N/A",
                redactor.sanitize_for_spreadsheet(ev.incident_id[:8] if ev.incident_id else "Unassigned"),
                redactor.sanitize_for_spreadsheet(file_name_map.get(ev.file_id, "unknown")),
                ev.original_line_number,
                redactor.sanitize_for_spreadsheet(ev.component or "N/A"),
                redactor.sanitize_for_spreadsheet(ev.message),
                redactor.sanitize_for_spreadsheet(ev.error_type or ev.log_level),
            ]

            for col_idx, val in enumerate(vals, start=1):
                cell = ws2.cell(row=row_idx, column=col_idx, value=val)
                cell.font = regular_font
                cell.border = thin_border
                cell.alignment = Alignment(vertical="top", wrap_text=(col_idx == 7))

        # -------------------------------------------------------------
        # Worksheet 3: Original Log Events
        # -------------------------------------------------------------
        ws3 = wb.create_sheet(title="Original Log Events")
        ws3.views.sheetView[0].showGridLines = True

        ws3["A1"] = "NoBugs AI — Original Parsed Log Events"
        ws3["A1"].font = title_font
        ws3["A2"] = "Complete audit record of sanitized source log events with preserved line numbers"
        ws3["A2"].font = sub_font

        headers3 = [
            "Source Filename",
            "Line Number",
            "Timestamp",
            "Sanitized Log Level",
            "Sanitized Message",
            "Associated Incident ID",
        ]

        row_idx = 4
        for col_idx, h in enumerate(headers3, start=1):
            cell = ws3.cell(row=row_idx, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for ev in events:
            row_idx += 1
            vals = [
                redactor.sanitize_for_spreadsheet(file_name_map.get(ev.file_id, "unknown")),
                ev.original_line_number,
                ev.timestamp.strftime("%Y-%m-%d %H:%M:%S") if ev.timestamp else "N/A",
                redactor.sanitize_for_spreadsheet(ev.log_level),
                redactor.sanitize_for_spreadsheet(ev.message),
                redactor.sanitize_for_spreadsheet(ev.incident_id[:8] if ev.incident_id else "None"),
            ]

            for col_idx, val in enumerate(vals, start=1):
                cell = ws3.cell(row=row_idx, column=col_idx, value=val)
                cell.font = regular_font
                cell.border = thin_border
                cell.alignment = Alignment(vertical="top", wrap_text=(col_idx == 5))

        # Adjust column widths for clean readability
        for ws in (ws1, ws2, ws3):
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.row < 4:
                        continue  # Skip title rows
                    val_str = str(cell.value or "")
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                # Cap width between 12 and 60
                ws.column_dimensions[col_letter].width = max(min(max_len + 3, 60), 12)

        # Output to buffer
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output


excel_service = ExcelReportService()
